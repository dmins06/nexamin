"""Convert the user-approved UTH review workbook to reproducible import data.

Usage:
  python scripts/normalize_uth_workbook.py INPUT.xlsx data/uth_curricula.json
"""

import argparse
import gzip
import hashlib
import json
import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook

EXPECTED_HEADERS = (
    "Ngành (PDF)", "Chuyên ngành/CTĐT (PDF)", "Mã ngành/CTĐT", "Khóa",
    "Ngành trên danh mục", "Chương trình trên danh mục", "STT tài liệu",
    "Mã môn", "Tên môn", "Số TC", "Bắt buộc/tự chọn", "Nhóm tự chọn",
    "Năm học", "Học kỳ", "Môn tiên quyết", "Ghi chú", "HP bắt buộc (thô)",
    "HP tự chọn (thô)", "Tổng TC CTĐT (văn bản)", "File nguồn", "Trang",
    "URL PDF", "Trang chương trình",
)


def clean(value: object) -> str | None:
    if value is None:
        return None
    text = " ".join(unicodedata.normalize("NFKC", str(value)).split())
    return text or None


def integer(value: object) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)) and int(value) == value and value > 0:
        return int(value)
    return None


def normalize(source: Path) -> dict[str, object]:
    book = load_workbook(source, read_only=True, data_only=True)
    sheet = book["Danh sách môn"]
    rows = sheet.iter_rows(values_only=True)
    for _ in range(5):
        next(rows)
    headers = tuple(next(rows))
    if headers != EXPECTED_HEADERS:
        raise ValueError("UTH workbook headers changed; inspect before importing")
    normalized = []
    for excel_row, values in enumerate(rows, start=7):
        if not any(value is not None for value in values):
            continue
        major = clean(values[0])
        specialization = clean(values[1])
        name = clean(values[8])
        source_file = clean(values[19])
        cohort_match = re.search(r"20\d{2}", clean(values[3]) or "")
        if not all((major, specialization, name, source_file, cohort_match)):
            raise ValueError(f"Incomplete UTH source row {excel_row}")
        group = clean(values[11])
        requirement = re.search(r"chọn\s*(\d+)\s*TC", group or "", flags=re.IGNORECASE)
        group_code = re.match(r"TC\d+", group or "", flags=re.IGNORECASE)
        normalized.append({
            "source_row": excel_row,
            "major": major,
            "specialization": specialization,
            "program_code": clean(values[2]) or "",
            "cohort": int(cohort_match.group()),
            "course_code": clean(values[7]),
            "course_name": name,
            "credits": integer(values[9]),
            "course_type": clean(values[10]) or "Không xác định",
            "group_label": group,
            "group_code": group_code.group().upper() if group_code else None,
            "group_required_credits": int(requirement.group(1)) if requirement else None,
            "academic_year": clean(values[12]),
            "semester": clean(values[13]),
            "prerequisite": clean(values[14]),
            "note": clean(values[15]),
            "source_file": source_file,
            "source_page": integer(values[20]),
            "source_url": clean(values[21]),
            "program_url": clean(values[22]),
        })
    if len(normalized) != 6971:
        raise ValueError(f"Expected 6971 UTH rows; found {len(normalized)}")
    return {
        "university_code": "UTH",
        "source_name": source.name,
        "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "approved": True,
        "rows": normalized,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    document = normalize(args.source)
    encoded = (json.dumps(document, ensure_ascii=False, separators=(",", ":")) + "\n").encode("utf-8")
    if args.output.suffix == ".gz":
        args.output.write_bytes(gzip.compress(encoded, compresslevel=9, mtime=0))
    else:
        args.output.write_bytes(encoded)
    print(f"Normalized {len(document['rows'])} UTH rows into {args.output}")


if __name__ == "__main__":
    main()
