"""Small synthetic archives/text only; no cloud, credentials or user files."""
import io
import uuid
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document as DocxDocument
from fastapi import UploadFile
from sqlalchemy import func, select

from app.core.exceptions import AppError
from app.models import Course, Document
from app.services.documents import DocumentService


class MemoryStorage:
    def __init__(self):
        self.objects: dict[str, bytes] = {}

    async def upload(self, key: str, content: bytes, content_type: str):
        self.objects[key] = content


@pytest.mark.asyncio
async def test_compressed_office_payload_rejected_before_blob_or_db_write(session, users):
    owner, _ = users
    course = Course(id=uuid.uuid4(), owner_id=owner.id, name="Resource fixture")
    session.add(course)
    await session.commit()
    source = DocxDocument()
    source.add_paragraph("A" * (2 * 1024 * 1024))
    output = io.BytesIO()
    source.save(output)
    content = output.getvalue()
    assert len(content) < 100 * 1024
    storage = MemoryStorage()
    with pytest.raises(AppError) as error:
        await DocumentService(session, owner.id, storage).upload(
            course.id, UploadFile(filename="synthetic.docx", file=io.BytesIO(content)),
        )
    assert error.value.status_code == 413
    assert storage.objects == {}
    assert await session.scalar(select(func.count()).select_from(Document)) == 0


def compressed_archive(name: str, data: bytes) -> bytes:
    stream = io.BytesIO()
    with ZipFile(stream, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr(name, data)
    return stream.getvalue()


def test_archive_expansion_and_member_limits_are_checked_before_office_parser():
    from app.workers.document_parser import DocumentRejected, ParserLimits, validate_archive

    content = compressed_archive("word/document.xml", b"x" * 8192)
    with pytest.raises(DocumentRejected) as error:
        validate_archive(content, ".docx", ParserLimits(max_archive_bytes=4096))
    assert error.value.code == "DOCUMENT_RESOURCE_LIMIT"
    with pytest.raises(DocumentRejected):
        validate_archive(content, ".docx", ParserLimits(max_member_bytes=4096))


@pytest.mark.parametrize("name", ["../word/document.xml", "/word/document.xml", "word\\document.xml"])
def test_unsafe_archive_names_are_rejected(name):
    from app.workers.document_parser import DocumentRejected, ParserLimits, validate_archive

    with pytest.raises(DocumentRejected) as error:
        validate_archive(compressed_archive(name, b"synthetic"), ".docx", ParserLimits())
    assert error.value.code == "INVALID_DOCUMENT"


def test_pdf_page_and_decompressed_stream_limits():
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, NameObject

    from app.workers.document_parser import DocumentRejected, ParserLimits, extract_pages

    writer = PdfWriter()
    page = writer.add_blank_page(width=200, height=200)
    stream = DecodedStreamObject()
    stream.set_data(b"% synthetic comment\n" * 128)
    page[NameObject("/Contents")] = stream.flate_encode()
    writer.add_blank_page(width=200, height=200)
    output = io.BytesIO()
    writer.write(output)
    for bounds in [ParserLimits(max_pages=1), ParserLimits(max_member_bytes=1024)]:
        with pytest.raises(DocumentRejected) as error:
            extract_pages(output.getvalue(), ".pdf", bounds)
        assert error.value.code == "DOCUMENT_RESOURCE_LIMIT"


@pytest.mark.asyncio
async def test_real_docx_pptx_utf8_are_parsed_in_child_without_blocking_loop():
    import asyncio

    from pptx import Presentation

    from app.services.document_processing import extract_in_child, processing_slot

    word = DocxDocument()
    word.add_paragraph("Tài liệu thử nghiệm tiếng Việt")
    word_bytes = io.BytesIO()
    word.save(word_bytes)
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[5])
    slide.shapes.title.text = "Nội dung slide riêng"
    slide_bytes = io.BytesIO()
    slides.save(slide_bytes)
    ticks = 0
    running = True

    async def heartbeat():
        nonlocal ticks
        while running:
            ticks += 1
            await asyncio.sleep(0.005)

    task = asyncio.create_task(heartbeat())
    try:
        async with processing_slot():
            assert await extract_in_child(word_bytes.getvalue(), ".docx") == [(None, "Tài liệu thử nghiệm tiếng Việt")]
            assert await extract_in_child(slide_bytes.getvalue(), ".pptx") == [(1, "Nội dung slide riêng")]
            assert await extract_in_child("Nội dung UTF-8".encode(), ".txt") == [(None, "Nội dung UTF-8")]
    finally:
        running = False
        await task
    assert ticks > 5


@pytest.mark.asyncio
async def test_parser_capacity_fails_fast_and_releases_slot(monkeypatch):
    from app.config import settings
    from app.services.document_processing import processing_slot

    monkeypatch.setattr(settings, "document_parser_concurrency", 1)
    async with processing_slot():
        with pytest.raises(AppError) as error:
            async with processing_slot():
                pytest.fail("Busy parser admitted a second document")
        assert error.value.code == "DOCUMENT_PROCESSING_BUSY"
    async with processing_slot():
        pass


@pytest.mark.asyncio
async def test_timeout_cancellation_output_limits_reap_child_and_scrub_secrets(tmp_path, monkeypatch):
    import asyncio

    from app.config import settings
    from app.services import document_processing as processing

    script = tmp_path / "synthetic_worker.py"
    script.write_text("import time\ntime.sleep(30)\n")
    monkeypatch.setattr(processing, "PARSER_SCRIPT", script)
    monkeypatch.setattr(settings, "document_parser_timeout_seconds", 0.1)
    monkeypatch.setenv("JWT_SECRET", "synthetic-must-not-reach-parser")
    monkeypatch.setenv("PRIVATE_BLOB_READ_WRITE_TOKEN", "synthetic-must-not-reach-parser")
    original = asyncio.create_subprocess_exec
    children = []
    started = asyncio.Event()

    async def spawn(*args, **kwargs):
        assert "JWT_SECRET" not in kwargs["env"]
        assert "PRIVATE_BLOB_READ_WRITE_TOKEN" not in kwargs["env"]
        assert "-I" in args
        child = await original(*args, **kwargs)
        children.append(child)
        started.set()
        return child

    monkeypatch.setattr(asyncio, "create_subprocess_exec", spawn)
    with pytest.raises(AppError) as error:
        await processing.extract_in_child(b"synthetic", ".txt")
    assert error.value.code == "DOCUMENT_RESOURCE_LIMIT"
    assert children[-1].returncode is not None
    started.clear()
    monkeypatch.setattr(settings, "document_parser_timeout_seconds", 10)
    task = asyncio.create_task(processing.extract_in_child(b"synthetic", ".txt"))
    await started.wait()
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task
    assert children[-1].returncode is not None
    script.write_text("import sys\nsys.stdout.write('x' * 16000)\n")
    monkeypatch.setattr(settings, "document_max_text_bytes", 1)
    monkeypatch.setattr(settings, "document_max_pages", 1)
    with pytest.raises(AppError) as error:
        await processing.extract_in_child(b"synthetic", ".txt")
    assert error.value.code == "DOCUMENT_RESOURCE_LIMIT"
    assert children[-1].returncode is not None


@pytest.mark.asyncio
@pytest.mark.parametrize("kind", ["memory", "cpu"])
async def test_real_os_memory_cpu_limits_are_enforced_in_child(kind, tmp_path, monkeypatch):
    from app.config import settings
    from app.services import document_processing as processing

    parser = str(processing.PARSER_SCRIPT)
    script = tmp_path / "os_limit_fixture.py"
    body = "return [(None, str(bytearray(limits.memory_mib * 2 * 1024 * 1024)))]" if kind == "memory" else "while True: pass"
    script.write_text(
        "import importlib.util,sys\n"
        f"spec=importlib.util.spec_from_file_location('synthetic_parser', {parser!r})\n"
        "module=importlib.util.module_from_spec(spec)\n"
        "sys.modules[spec.name]=module\n"
        "spec.loader.exec_module(module)\n"
        f"def exhaust(content,extension,limits):\n    {body}\n"
        "module.extract_pages=exhaust\nmodule.main()\n",
    )
    monkeypatch.setattr(processing, "PARSER_SCRIPT", script)
    monkeypatch.setattr(settings, "document_parser_cpu_seconds", 1)
    monkeypatch.setattr(settings, "document_parser_timeout_seconds", 5)
    with pytest.raises(AppError) as error:
        await processing.extract_in_child(b"synthetic", ".txt")
    assert error.value.code == "DOCUMENT_RESOURCE_LIMIT"


def test_chunk_budget_stops_before_large_embedding_allocation():
    from app.ai.rag import TextSplitter

    with pytest.raises(ValueError, match="maximum chunk"):
        TextSplitter(chunk_size=8, overlap=1).split([(1, "abcdef " * 20)], max_chunks=2)


@pytest.mark.asyncio
async def test_child_start_failure_is_safe_503(monkeypatch):
    import asyncio

    from app.services.document_processing import extract_in_child

    async def unavailable(*args, **kwargs):
        raise OSError("synthetic OS detail must not reach client")

    monkeypatch.setattr(asyncio, "create_subprocess_exec", unavailable)
    with pytest.raises(AppError) as error:
        await extract_in_child(b"synthetic", ".txt")
    assert error.value.status_code == 503
    assert "synthetic OS detail" not in error.value.message


@pytest.mark.parametrize("payload", [
    b'{"pages":[[true,"text"]]}', b'{"pages":[[0,"text"]]}',
    b'{"pages":[[null,"\\ud800"]]}', b'{"error":"supplied secret detail"}', b'[]',
])
def test_bad_child_contract_cannot_leak_details(payload):
    from app.services.document_processing import decode_pages
    from app.workers.document_parser import ParserLimits

    with pytest.raises(AppError) as error:
        decode_pages(payload, ParserLimits())
    assert error.value.code == "DOCUMENT_PROCESSOR_UNAVAILABLE"
    assert "supplied secret" not in error.value.message


@pytest.mark.asyncio
async def test_upload_api_preflight_processing_and_private_download_contract(session, users, monkeypatch):
    from contextlib import asynccontextmanager

    import httpx
    from fastapi import FastAPI

    from app.api.v1 import documents
    from app.config import settings
    from app.core.exceptions import register_exception_handlers
    from app.database import get_db
    from app.services.auth import AuthService
    from app.workers import document_tasks

    owner, other = users
    course = Course(id=uuid.uuid4(), owner_id=owner.id, name="API fixture")
    session.add(course)
    await session.commit()
    storage = MemoryStorage()

    async def download(key):
        return storage.objects[key]

    storage.download = download
    monkeypatch.setattr(documents, "document_storage", lambda: storage)
    monkeypatch.setattr(settings, "storage_provider", "blob")

    async def get_session():
        yield session

    @asynccontextmanager
    async def factory():
        yield session

    monkeypatch.setattr(document_tasks, "SessionFactory", factory)
    app = FastAPI()
    app.include_router(documents.router, prefix="/api/v1")
    app.dependency_overrides[get_db] = get_session
    register_exception_handlers(app)
    pair = await AuthService(session)._issue_pair(owner)
    other_pair = await AuthService(session)._issue_pair(other)
    headers = {"Authorization": "Bearer " + pair.access_token}
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://localhost") as client:
        content = "Tài liệu API synthetic".encode()
        response = await client.post("/api/v1/documents/upload", headers=headers, data={"course_id": str(course.id)}, files={"file": ("fixture.txt", content, "text/plain")})
        assert response.status_code == 202
        assert response.json()["processing_status"] == "ready"
        docid = response.json()["id"]
        response = await client.get(f"/api/v1/documents/{docid}/download", headers=headers)
        assert response.content == content
        assert response.headers["cache-control"] == "private, no-store"
        denied = await client.get(f"/api/v1/documents/{docid}/download", headers={"Authorization": "Bearer " + other_pair.access_token})
        assert denied.status_code == 404
        invalid = await client.post("/api/v1/documents/upload", headers=headers, data={"course_id": str(course.id)}, files={"file": ("invalid.pdf", b"%PDF-1.7\nsynthetic invalid\n", "application/pdf")})
        assert invalid.status_code == 422
        assert invalid.json()["error"]["code"] == "INVALID_DOCUMENT"
    assert len(storage.objects) == 1
    assert await session.scalar(select(func.count()).select_from(Document)) == 1
