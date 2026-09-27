"""Offline-first extraction of review questions from rights-approved documents.

Discovery and rights review happen before this module is called. Parsed questions
remain pending and never become a student exam without a separate answer review.
"""

import csv
import io
import ipaddress
import json
import re
import socket
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import urlsplit
from urllib.robotparser import RobotFileParser

import httpx
from docx import Document
from openpyxl import load_workbook  # type: ignore[import-untyped]
from pypdf import PdfReader

from app.models.practice import QuestionSource
from app.services.practice_import import ImportedOption, ImportedQuestion

MAX_SOURCE_BYTES = 8 * 1024 * 1024
USER_AGENT = "NexaminPracticeImporter/1.0 (+https://nexamin.id.vn)"


class _TextFromHTML(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self.hidden = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style", "nav", "footer"}:
            self.hidden += 1
        if tag in {"p", "div", "br", "li", "h1", "h2", "h3", "h4"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style", "nav", "footer"} and self.hidden:
            self.hidden -= 1
        if tag in {"p", "div", "li", "h1", "h2", "h3", "h4"}:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self.hidden:
            self.parts.append(data)


def extract_text(data: bytes, filename: str) -> str:
    """Read supported public file formats without executing embedded content."""
    suffix = Path(filename).suffix.lower()
    if len(data) > MAX_SOURCE_BYTES:
        raise ValueError("Source file exceeds 8 MB.")
    if suffix in {".txt", ".md"}:
        return data.decode("utf-8-sig")
    if suffix in {".html", ".htm"}:
        parser = _TextFromHTML()
        parser.feed(data.decode("utf-8-sig", errors="replace"))
        return "".join(parser.parts)
    if suffix == ".pdf":
        return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
    if suffix == ".docx":
        return "\n".join(paragraph.text for paragraph in Document(io.BytesIO(data)).paragraphs)
    if suffix == ".csv":
        return "\n".join(" | ".join(row) for row in csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    if suffix == ".xlsx":
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            return "\n".join(
                " | ".join(str(cell) if cell is not None else "" for cell in row)
                for sheet in workbook.worksheets for row in sheet.values
            )
        finally:
            workbook.close()
    if suffix == ".json":
        document = json.loads(data)
        return json.dumps(document, ensure_ascii=False, indent=2)
    raise ValueError(f"Unsupported source format: {suffix}")


_QUESTION = re.compile(r"(?im)^\s*(?:câu\s*)?(\d{1,3})\s*[.):]\s*(.+)$")
_OPTION = re.compile(r"(?im)^\s*([A-H])\s*[.):]\s*(.+)$")
_ANSWER = re.compile(r"(?im)^\s*(?:đáp\s*án|answer)\s*[:：]\s*([A-H])\b")


@dataclass(frozen=True)
class ExtractionResult:
    questions: list[ImportedQuestion]
    candidate_count: int
    rejected_count: int


def extract_questions_detailed(text: str) -> ExtractionResult:
    """Conservative extractor; unanswered or uncertain items stay pending."""
    review_heading = re.search(r"(?i)câu\s+hỏi\s+(?:(?:thảo\s+luận|thao\s+luan)\s+và\s+)?ôn\s+tập", text)
    if review_heading:
        text = text[review_heading.end():]
    markers = list(_QUESTION.finditer(text))
    results: list[ImportedQuestion] = []
    for index, marker in enumerate(markers):
        block = text[marker.start():markers[index + 1].start() if index + 1 < len(markers) else len(text)]
        lines = [" ".join(line.split()) for line in block.splitlines() if line.strip()]
        if not lines:
            continue
        first = _QUESTION.match(lines[0])
        if first is None:
            continue
        question_text = first.group(2).strip()
        if len(question_text) < 10:
            continue
        options = [(match.group(1), match.group(2).strip()) for line in lines[1:]
                   if (match := _OPTION.match(line))]
        answer = next((match.group(1) for line in lines[1:]
                       if (match := _ANSWER.match(line))), None)
        if options and (len(options) < 2 or len({key for key, _ in options}) != len(options)):
            continue
        if options and answer is not None and answer not in {key for key, _ in options}:
            continue
        if not options and "?" not in question_text and not re.match(
            r"(?i)^(phân tích|trình bày|nêu|giải thích|so sánh|tại sao|vì sao)", question_text
        ):
            continue
        results.append(ImportedQuestion(
            external_id=first.group(1),
            text=question_text,
            type="single_choice" if options else "short_answer",
            difficulty="medium",
            options=[ImportedOption(key=key, text=value, correct=key == answer)
                     for key, value in options],
        ))
    return ExtractionResult(results, len(markers), len(markers) - len(results))


def extract_questions(text: str) -> list[ImportedQuestion]:
    return extract_questions_detailed(text).questions


def _public_host(host: str) -> bool:
    if host == "localhost" or host.endswith(".local"):
        return False
    try:
        addresses = socket.getaddrinfo(host, 443, type=socket.SOCK_STREAM)
    except OSError:
        return False
    return bool(addresses) and all(
        ipaddress.ip_address(address[4][0]).is_global for address in addresses
    )


async def fetch_approved(source: QuestionSource) -> tuple[bytes, str]:
    """Fetch one approved HTTPS resource, respecting robots and strict size limits."""
    parsed = urlsplit(source.url)
    if (
        source.import_status not in {"approved", "imported"} or not source.rights_url
        or parsed.scheme != "https" or not parsed.hostname or not _public_host(parsed.hostname)
    ):
        raise ValueError("Source is not approved or URL is not a public HTTPS host.")
    origin = f"https://{parsed.hostname}"
    async with httpx.AsyncClient(timeout=20, follow_redirects=False, headers={"User-Agent": USER_AGENT}) as client:
        robots_response = await client.get(f"{origin}/robots.txt")
        if robots_response.status_code == 200:
            robots = RobotFileParser()
            robots.parse(robots_response.text.splitlines())
            if not robots.can_fetch(USER_AGENT, source.url):
                raise ValueError("Source robots.txt disallows fetching this URL.")
        elif robots_response.status_code != 404:
            raise ValueError("Cannot verify source robots.txt.")
        async with client.stream("GET", source.url) as response:
            response.raise_for_status()
            chunks: list[bytes] = []
            total = 0
            async for chunk in response.aiter_bytes():
                total += len(chunk)
                if total > MAX_SOURCE_BYTES:
                    raise ValueError("Source response exceeds 8 MB.")
                chunks.append(chunk)
            content_type = response.headers.get("content-type", "").split(";")[0].lower()
            suffix = {
                "text/html": ".html", "application/pdf": ".pdf",
                "text/plain": ".txt", "text/csv": ".csv",
                "application/json": ".json",
            }.get(content_type, Path(parsed.path).suffix.lower())
            return b"".join(chunks), suffix
