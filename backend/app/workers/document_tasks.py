import asyncio
import io
import uuid
from pathlib import Path

from docx import Document as DocxDocument
from pptx import Presentation
from pypdf import PdfReader
from sqlalchemy import delete

from app.ai.rag import TextSplitter
from app.config import settings
from app.core.storage import selected_storage
from app.database import SessionFactory
from app.models import Document, DocumentChunk, ProcessingStatus
from app.workers.celery_app import celery_app


def extract_pages(content: bytes, extension: str) -> list[tuple[int | None, str]]:
    if extension == ".pdf":
        reader = PdfReader(io.BytesIO(content))
        return [(index + 1, page.extract_text() or "") for index, page in enumerate(reader.pages)]
    if extension == ".docx":
        document = DocxDocument(io.BytesIO(content))
        return [(None, "\n".join(paragraph.text for paragraph in document.paragraphs))]
    if extension == ".pptx":
        presentation = Presentation(io.BytesIO(content))
        pages: list[tuple[int | None, str]] = []
        for index, slide in enumerate(presentation.slides):
            texts = [shape.text for shape in slide.shapes if hasattr(shape, "text")]
            pages.append((index + 1, "\n".join(texts)))
        return pages
    return [(None, content.decode("utf-8"))]


async def process_document(document_id: uuid.UUID) -> None:
    async with SessionFactory() as session:
        document = await session.get(Document, document_id)
        if document is None:
            return
        document.processing_status = ProcessingStatus.PROCESSING
        document.processing_error = None
        await session.commit()
        try:
            backend = selected_storage()
            content = await backend.download(document.storage_key)
            pages = extract_pages(content, Path(document.filename).suffix.lower())
            pieces = TextSplitter().split(pages)
            embeddings = [[0.0] * settings.embedding_dimensions for _ in pieces]
            await session.execute(
                delete(DocumentChunk).where(DocumentChunk.document_id == document.id)
            )
            for index, ((page_number, text), embedding) in enumerate(
                zip(pieces, embeddings, strict=True)
            ):
                session.add(
                    DocumentChunk(
                        document_id=document.id,
                        course_id=document.course_id,
                        content=text,
                        page_number=page_number,
                        chunk_index=index,
                        embedding=embedding,
                        chunk_metadata={
                            "source": document.original_filename,
                            "embedding_provider": "local",
                        },
                    )
                )
            document.page_count = len(pages)
            document.processing_status = ProcessingStatus.READY
            await session.commit()
        except Exception as exc:
            await session.rollback()
            failed = await session.get(Document, document_id)
            if failed is not None:
                failed.processing_status = ProcessingStatus.FAILED
                failed.processing_error = str(exc)[:1000]
                await session.commit()
            raise


@celery_app.task(  # type: ignore[untyped-decorator]
    name="documents.process",
    autoretry_for=(ConnectionError, TimeoutError),
    retry_backoff=True,
    retry_kwargs={"max_retries": 3},
)
def process_document_task(document_id: str) -> None:
    asyncio.run(process_document(uuid.UUID(document_id)))
