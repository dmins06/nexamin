"""Bounded, in-memory parsing. Also runs as an isolated Python child script.

No app configuration, database, storage client or credentials are imported here.
"""
import io
import json
import math
import sys
from dataclasses import dataclass
from pathlib import PurePosixPath
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile


@dataclass(frozen=True, slots=True)
class ParserLimits:
    max_input_bytes: int = 25 * 1024 * 1024
    max_archive_bytes: int = 48 * 1024 * 1024
    max_member_bytes: int = 16 * 1024 * 1024
    max_archive_entries: int = 2048
    max_compression_ratio: int = 200
    max_pages: int = 500
    max_text_bytes: int = 1024 * 1024
    memory_mib: int = 256
    cpu_seconds: int = 15

    def arguments(self) -> list[str]:
        return [str(value) for value in (
            self.max_input_bytes, self.max_archive_bytes, self.max_member_bytes,
            self.max_archive_entries, self.max_compression_ratio, self.max_pages,
            self.max_text_bytes, self.memory_mib, self.cpu_seconds,
        )]


class DocumentRejected(ValueError):
    def __init__(self, code: str):
        self.code = code
        super().__init__(code)


def validate_archive(content: bytes, extension: str, limits: ParserLimits) -> None:
    # ZipFile allocates the central directory before exposing infolist(). Limit
    # that byte area independently, including forged entry counts and ZIP64.
    tail_start = max(0, len(content) - 65557)
    end = content.rfind(b"PK\x05\x06", tail_start)
    if end < 0 or len(content) - end < 22:
        raise DocumentRejected("INVALID_DOCUMENT")
    if end >= 20 and content[end - 20:end - 16] == b"PK\x06\x07":
        raise DocumentRejected("DOCUMENT_RESOURCE_LIMIT")
    directory_bytes = int.from_bytes(content[end + 12:end + 16], "little")
    entry_count = int.from_bytes(content[end + 10:end + 12], "little")
    if directory_bytes > 1024 * 1024 or entry_count > limits.max_archive_entries:
        raise DocumentRejected("DOCUMENT_RESOURCE_LIMIT")
    with ZipFile(io.BytesIO(content)) as archive:
        entries = archive.infolist()
        if len(entries) > limits.max_archive_entries:
            raise DocumentRejected("DOCUMENT_RESOURCE_LIMIT")
        names: set[str] = set()
        total = 0
        for entry in entries:
            path = PurePosixPath(entry.filename)
            if (entry.filename in names or path.is_absolute() or ".." in path.parts
                    or "\\" in entry.filename or "\x00" in entry.filename):
                raise DocumentRejected("INVALID_DOCUMENT")
            names.add(entry.filename)
            if entry.flag_bits & 1 or entry.compress_type not in (ZIP_STORED, ZIP_DEFLATED):
                raise DocumentRejected("INVALID_DOCUMENT")
            total += entry.file_size
            if (entry.file_size > limits.max_member_bytes or total > limits.max_archive_bytes
                    or entry.file_size > max(entry.compress_size, 1) * limits.max_compression_ratio):
                raise DocumentRejected("DOCUMENT_RESOURCE_LIMIT")
        required = {"[Content_Types].xml", "_rels/.rels",
                    "word/document.xml" if extension == ".docx" else "ppt/presentation.xml"}
        if not required.issubset(names):
            raise DocumentRejected("INVALID_DOCUMENT")
        # Verify actual decompression/CRC too, with bounded reads and counters.
        actual_total = 0
        for entry in entries:
            actual_member = 0
            with archive.open(entry) as member:
                while block := member.read(64 * 1024):
                    actual_member += len(block)
                    actual_total += len(block)
                    if actual_member > limits.max_member_bytes or actual_total > limits.max_archive_bytes:
                        raise DocumentRejected("DOCUMENT_RESOURCE_LIMIT")


def extract_pages(
    content: bytes, extension: str, limits: ParserLimits | None = None,
) -> list[tuple[int | None, str]]:
    bounds = limits or ParserLimits()
    if len(content) > bounds.max_input_bytes:
        raise DocumentRejected("DOCUMENT_RESOURCE_LIMIT")
    pages: list[tuple[int | None, str]] = []
    text_bytes = 0

    def append(number: int | None, text: str) -> None:
        nonlocal text_bytes
        text_bytes += len(text.encode("utf-8"))
        if len(pages) >= bounds.max_pages or text_bytes > bounds.max_text_bytes:
            raise DocumentRejected("DOCUMENT_RESOURCE_LIMIT")
        pages.append((number, text))

    if extension == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            raise DocumentRejected("INVALID_DOCUMENT")
        if len(reader.pages) > bounds.max_pages:
            raise DocumentRejected("DOCUMENT_RESOURCE_LIMIT")
        for index, page in enumerate(reader.pages):
            stream = page.get_contents()
            if stream is not None and len(stream.get_data()) > bounds.max_member_bytes:
                raise DocumentRejected("DOCUMENT_RESOURCE_LIMIT")
            append(index + 1, page.extract_text() or "")
    elif extension in (".docx", ".pptx"):
        validate_archive(content, extension, bounds)
        if extension == ".docx":
            from docx import Document

            document = Document(io.BytesIO(content))
            parts: list[str] = []
            size = 0
            for paragraph in document.paragraphs:
                text = paragraph.text
                size += len(text.encode("utf-8")) + (1 if parts else 0)
                if size > bounds.max_text_bytes:
                    raise DocumentRejected("DOCUMENT_RESOURCE_LIMIT")
                parts.append(text)
            append(None, "\n".join(parts))
        else:
            from pptx import Presentation

            presentation = Presentation(io.BytesIO(content))
            if len(presentation.slides) > bounds.max_pages:
                raise DocumentRejected("DOCUMENT_RESOURCE_LIMIT")
            for index, slide in enumerate(presentation.slides):
                text = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
                append(index + 1, text)
    elif extension == ".txt":
        append(None, content.decode("utf-8"))
    else:
        raise DocumentRejected("INVALID_DOCUMENT")
    return pages


def main() -> None:
    try:
        values = [int(value) for value in sys.argv[2:]]
        limits = ParserLimits(*values)
        import resource

        resource.setrlimit(resource.RLIMIT_AS, (limits.memory_mib * 1024 * 1024,) * 2)
        resource.setrlimit(resource.RLIMIT_CPU, (math.ceil(limits.cpu_seconds), limits.cpu_seconds + 1))
        resource.setrlimit(resource.RLIMIT_CORE, (0, 0))
        resource.setrlimit(resource.RLIMIT_FSIZE, (0, 0))
        content = sys.stdin.buffer.read(limits.max_input_bytes + 1)
        result: dict[str, object] = {"pages": extract_pages(content, sys.argv[1], limits)}
    except DocumentRejected as error:
        result = {"error": error.code}
    except MemoryError:
        result = {"error": "DOCUMENT_RESOURCE_LIMIT"}
    except (ImportError, OSError):
        result = {"error": "DOCUMENT_PROCESSOR_UNAVAILABLE"}
    except Exception:
        # Deliberately exclude parser messages and supplied file contents.
        result = {"error": "INVALID_DOCUMENT"}
    sys.stdout.buffer.write(json.dumps(result, ensure_ascii=False).encode("utf-8"))


if __name__ == "__main__":
    main()
